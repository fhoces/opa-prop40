"""Python twin of R/revenue_chain.R. Reads the same workbook (cached values,
data_only=True) and re-derives the same per-row formula: O = (D - K) * rate.
Deterministic - must match R to 1e-9 (modulo openpyxl vs readxl float
parsing, which is well under that)."""
import warnings

import openpyxl


def _read_calculations_sheet(path, sheet_name, first_row=3, last_row=214):
    warnings.filterwarnings("ignore")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name]
    net_worth, trre, moving = [], [], []
    for row in range(first_row, last_row + 1):
        net_worth.append(ws[f"D{row}"].value)
        trre.append(ws[f"K{row}"].value)
        moving.append(ws[f"F{row}"].value)
    return {"net_worth": net_worth, "trre": trre, "moving": moving}


def extract_calculations_preferred(path):
    return _read_calculations_sheet(path, "Calculations_Preferred")


def extract_calculations_expanded(path):
    return _read_calculations_sheet(path, "Calculations_Expanded")


def _row_wt(sheet, tax_rate=0.05):
    return [(d - (k or 0)) * tax_rate for d, k in zip(sheet["net_worth"], sheet["trre"])]


def compute_revenue_chain(calc_preferred, calc_expanded, baseline_literal=94.2,
                           brulhart_semi_elasticity=10.32, tax_rate=0.05):
    wt_preferred = _row_wt(calc_preferred, tax_rate)
    wt_expanded = _row_wt(calc_expanded, tax_rate)

    baseline_precise = sum(wt_preferred) / 1e9
    confirmed6_ceiling = sum(
        v for v, m in zip(wt_preferred, calc_preferred["moving"]) if m == "N"
    ) / 1e9
    expanded10_estimate = sum(
        v for v, m in zip(wt_expanded, calc_expanded["moving"]) if m == "N"
    ) / 1e9

    migration_share = brulhart_semi_elasticity * tax_rate
    literature_calibrated = (1 - migration_share) * baseline_literal

    return {
        "baseline_precise": baseline_precise,
        "baseline_literal": baseline_literal,
        "confirmed6_ceiling": confirmed6_ceiling,
        "expanded10_estimate": expanded10_estimate,
        "migration_share": migration_share,
        "literature_calibrated": literature_calibrated,
    }
