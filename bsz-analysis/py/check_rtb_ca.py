"""Check query 1 (sql/01_rtb_ca.sql) against its answer keys.

Run from bsz-analysis/, after py/run_sql.py and R/run_sql.R have exported
the three tables:

    python py/check_rtb_ca.py

Writes data-raw/sql-out/check_rtb_ca.md (gitignored) and prints a summary.
Exit code 1 if any check fails, 0 otherwise.

Answer keys:
  * private: the authors' private sheets and their own export of the daily
    aggregate, located through data-raw/private-paths.csv (gitignored; see
    data-raw/private-paths.example.csv);
  * public: BSZ_MainTablesFigures.xlsx, sheets rtb_2026_industry and
    data_sec_agg (whose year rows feed shortrunseries columns C and Q).

Privacy rule for the report: rows of the private sheets are never printed,
only counts, row totals and maximum absolute differences. Rows of the public
workbook may be printed.
"""
import csv
import math
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from bundle_paths import private_file, public_workbook_path, sql_out_dir  # noqa: E402

TOL = 1e-6              # recomputed vs answer key
TOL_PARITY = 1e-9       # R export vs Python export
TOL_2DP = 0.005 + 1e-9  # data_sec_agg holds values typed at 2 decimals

EOY_DATES = ["2019-12-31", "2020-12-31", "2021-12-31", "2022-12-31",
             "2023-12-31", "2024-12-31", "2026-01-01"]
EOY_NUM = ["forbes_worth", "forbes_public_worth", "forbes_private_worth"]
EOY_TEXT = ["forbes_name", "state", "country_citizenship", "source", "industries"]
IND_COLS = ["n_billionaires", "forbes_public_worth", "forbes_worth",
            "fraction_public_worth", "fraction_forbes_worth"]
AGG_NUM = ["forbes_worth_total", "forbes_worth_top4"]

TABLES = ["rtb_ca_eoy", "rtb_ca_2026_01_01_industry", "rtb_ca_aggregate"]


# ---------------------------------------------------------------------------
# readers
# ---------------------------------------------------------------------------
def _missing(v):
    return v is None or v == "" or v == "#N/A" or (isinstance(v, float) and math.isnan(v))


def _num(v):
    return None if _missing(v) else float(v)


def _iso(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return str(v)[:10]


def read_export(path):
    """A CSV written by run_sql (either language) as a list of dicts."""
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def read_sheet(wb, name, header_row=1, ncols=None):
    """Rows of a sheet as dicts keyed by the header; empty rows dropped."""
    rows = wb[name].iter_rows(min_row=header_row, values_only=True)
    header = list(next(rows))
    if ncols:
        header = header[:ncols]
    keep = [i for i, h in enumerate(header) if h is not None]
    out = []
    for r in rows:
        if all(v is None for v in r):
            continue
        out.append({header[i]: (r[i] if i < len(r) else None) for i in keep})
    return out


# ---------------------------------------------------------------------------
# comparisons; each returns a dict with ok, n, maxdiff and a note
# ---------------------------------------------------------------------------
def cmp_eoy(eoy, key_wb):
    out = []
    for d in EOY_DATES:
        ours = {r["forbes_id"]: r for r in eoy if r["date"] == d}
        key = {r["forbes_id"]: r for r in read_sheet(key_wb, "rtb_ca_" + d.replace("-", "_"))}
        common = sorted(set(ours) & set(key))
        maxd, text_bad, cik_bad = 0.0, 0, 0
        for i in common:
            for c in EOY_NUM:
                a, b = _num(ours[i][c]), _num(key[i][c])
                if (a is None) != (b is None):
                    maxd = math.inf
                elif a is not None:
                    maxd = max(maxd, abs(a - b))
            for c in EOY_TEXT:
                a, b = ours[i][c], key[i][c]
                if not ((_missing(a) and _missing(b)) or str(a) == str(b)):
                    text_bad += 1
            if d == "2026-01-01":
                if _num(ours[i]["sec_cik"]) != _num(key[i].get("sec_cik")):
                    cik_bad += 1
        extra, missing = len(set(ours) - set(key)), len(set(key) - set(ours))
        ok = (len(ours) == len(key) and extra == 0 and missing == 0
              and maxd <= TOL and text_bad == 0 and cik_bad == 0)
        note = (f"rows {len(ours)} vs {len(key)}; ids only in ours {extra}, only in key {missing}; "
                f"text mismatches {text_bad}" + (f"; sec_cik mismatches {cik_bad}" if d == "2026-01-01" else ""))
        out.append({"name": f"year-end list {d}", "ok": ok, "n": len(common), "maxdiff": maxd, "note": note})
    return out


def _cmp_rows(ours, key, cols, tol, label):
    """Align two lists of dicts on 'industries'; returns check + per-row lines."""
    o = {r["industries"]: r for r in ours}
    k = {r["industries"]: r for r in key}
    maxd, lines, nbad = 0.0, [], 0
    for name in sorted(set(o) & set(k)):
        for c in cols:
            a, b = _num(o[name][c]), _num(k[name][c])
            d = math.inf if (a is None) != (b is None) else (0.0 if a is None else abs(a - b))
            exact = c == "n_billionaires"
            bad = d != 0 if exact else d > tol
            if not exact:
                maxd = max(maxd, d)
            if bad:
                nbad += 1
                lines.append(f"| {name} | {c} | {a!r} | {b!r} | {d:.3g} |")
    extra, missing = len(set(o) - set(k)), len(set(k) - set(o))
    ok = nbad == 0 and extra == 0 and missing == 0
    note = f"rows {len(o)} vs {len(k)}; only in ours {extra}, only in key {missing}; cells off {nbad}"
    return {"name": label, "ok": ok, "n": len(set(o) & set(k)), "maxdiff": maxd, "note": note}, lines


def cmp_industry_private(ind, key_wb):
    key = read_sheet(key_wb, "rtb_ca_2026_01_01_industry", ncols=6)
    chk, _ = _cmp_rows(ind, key, IND_COLS, TOL, "industry table vs private sheet")
    return chk


def cmp_industry_public(ind, pub_wb):
    """Public rtb_2026_industry: 'Finance' is our 'Finance & Investments';
    its hand-added 'Other' subtotal row has no counterpart and is skipped."""
    key = read_sheet(pub_wb, "rtb_2026_industry", header_row=4, ncols=6)
    cut = next((i for i, r in enumerate(key) if r["industries"] == "Total"), len(key))
    key = [dict(r) for r in key[: cut + 1] if r["industries"] != "Other"]
    for r in key:
        if r["industries"] == "Finance":
            r["industries"] = "Finance & Investments"
    return _cmp_rows(ind, key, IND_COLS, TOL, "industry table vs public rtb_2026_industry")


def cmp_aggregate(agg, key_rows, label):
    o = {r["date"]: r for r in agg}
    k, dup_rows, dup_differ = {}, 0, 0
    for r in key_rows:
        d = _iso(r["date"])
        if d in k:
            # A date listed twice in the key: compare the copies, keep one.
            dup_rows += 1
            dup_differ += any(k[d][c] != r[c] for c in r if c != "date")
        else:
            k[d] = r
    common = sorted(set(o) & set(k))
    maxd, n_bad, dates_bad = 0.0, 0, 0
    for d in common:
        bad = int(o[d]["n_billionaires"]) != int(k[d]["n_billionaires"])
        n_bad += bad
        for c in AGG_NUM:
            diff = abs(float(o[d][c]) - float(k[d][c]))
            maxd = max(maxd, diff)
            bad = bad or diff > TOL
        dates_bad += bad
    only_o, only_k = sorted(set(o) - set(k)), sorted(set(k) - set(o))
    ok = dates_bad == 0 and not only_k and dup_differ == 0
    note = (f"dates {len(o)} vs {len(k)}, compared {len(common)}; dates off {dates_bad} "
            f"(count off on {n_bad}); only in ours {len(only_o)}, only in key {len(only_k)}"
            + (f"; key repeats {dup_rows} rows ({dup_differ} not identical)" if dup_rows else ""))
    return {"name": label, "ok": ok, "n": len(common), "maxdiff": maxd, "note": note,
            "only_ours": only_o, "only_key": only_k}


def cmp_data_sec_agg(agg, pub_wb):
    """Public data_sec_agg year rows (feed shortrunseries C and Q) vs our year-end totals."""
    by_year = {str(r["year"]): r for r in read_sheet(pub_wb, "data_sec_agg", header_row=4)}
    o = {r["date"]: r for r in agg}
    lines, maxd, ok = [], 0.0, True
    for d in EOY_DATES:
        year = "2025" if d == "2026-01-01" else d[:4]
        k = by_year[year]
        n_o, w_o = int(o[d]["n_billionaires"]), float(o[d]["forbes_worth_total"])
        dw = abs(w_o - k["forbes_worth"])
        maxd = max(maxd, dw)
        row_ok = n_o == k["n"] and dw <= TOL_2DP
        ok = ok and row_ok
        lines.append(f"| {year} | {n_o} | {k['n']} | {w_o:.4f} | {k['forbes_worth']} | "
                     f"{'ok' if row_ok else 'OFF'} |")
    chk = {"name": "year-end totals vs public data_sec_agg", "ok": ok, "n": len(EOY_DATES),
           "maxdiff": maxd, "note": "n exact; forbes_worth within 0.005 (2-decimal cells)"}
    return chk, lines


def cmp_parity():
    out = []
    for t in TABLES:
        p, r = sql_out_dir("py") / f"{t}.csv", sql_out_dir("r") / f"{t}.csv"
        if not r.exists():
            out.append({"name": f"R vs Python: {t}", "ok": False, "n": 0, "maxdiff": math.nan,
                        "note": "R export missing; run Rscript R/run_sql.R first"})
            continue
        a, b = read_export(p), read_export(r)
        same_cols = (list(a[0]) == list(b[0])) if a and b else len(a) == len(b)
        maxd, text_bad = 0.0, 0
        for x, y in zip(a, b):
            for c in x:
                u, v = x[c], y.get(c)
                try:
                    maxd = max(maxd, abs(float(u) - float(v)))
                except (TypeError, ValueError):
                    text_bad += u != v
        ok = same_cols and len(a) == len(b) and maxd <= TOL_PARITY and text_bad == 0
        out.append({"name": f"R vs Python: {t}", "ok": ok, "n": len(a), "maxdiff": maxd,
                    "note": f"rows {len(a)} vs {len(b)}; text cells off {text_bad}"})
    return out


def _fmt(x):
    return "inf" if x == math.inf else ("n/a" if isinstance(x, float) and math.isnan(x) else f"{x:.3g}")


def main():
    import openpyxl

    py = sql_out_dir("py")
    eoy = read_export(py / "rtb_ca_eoy.csv")
    ind = read_export(py / "rtb_ca_2026_01_01_industry.csv")
    agg = read_export(py / "rtb_ca_aggregate.csv")
    key_wb = openpyxl.load_workbook(private_file("private_sheets"), read_only=True, data_only=True)
    exp_wb = openpyxl.load_workbook(private_file("private_aggregate_export"),
                                    read_only=True, data_only=True)
    pub_wb = openpyxl.load_workbook(public_workbook_path(), read_only=True, data_only=True)

    checks = cmp_eoy(eoy, key_wb)
    checks.append(cmp_industry_private(ind, key_wb))
    pub_ind, pub_lines = cmp_industry_public(ind, pub_wb)
    checks.append(pub_ind)
    agg_key = cmp_aggregate(agg, read_sheet(key_wb, "rtb_ca_aggregate"), "daily aggregate vs private sheet")
    agg_exp = cmp_aggregate(agg, read_sheet(exp_wb, exp_wb.sheetnames[0]),
                            "daily aggregate vs the authors' aggregate export")
    checks += [agg_key, agg_exp]
    dsa, dsa_lines = cmp_data_sec_agg(agg, pub_wb)
    checks.append(dsa)
    checks += cmp_parity()
    failures = sum(not c["ok"] for c in checks)

    md = ["# Query 1 check report (sql/01_rtb_ca.sql)", "",
          f"Generated {datetime.now():%Y-%m-%d %H:%M}. Counts, totals and maximum absolute",
          "differences only; no row of the private sheets is printed.", "",
          f"Tolerances: {TOL:g} against answer keys, {TOL_PARITY:g} for R vs Python, 0.005 for the",
          "2-decimal cells of data_sec_agg. Counts must match exactly.", "",
          "## Summary", "",
          "| Check | Compared | Max abs diff | Status | Detail |",
          "|---|---:|---:|---|---|"]
    for c in checks:
        md.append(f"| {c['name']} | {c['n']} | {_fmt(c['maxdiff'])} | "
                  f"{'PASS' if c['ok'] else 'FAIL'} | {c['note']} |")
    md += ["", "## Daily aggregate: dates on one side only", "",
           f"Only in ours ({len(agg_key['only_ours'])}): " + (", ".join(agg_key["only_ours"]) or "none"), "",
           f"Only in the private sheet ({len(agg_key['only_key'])}): "
           + (", ".join(agg_key["only_key"]) or "none"), "",
           "forbes_private_worth is in neither answer key, so only R vs Python parity checks it.", "",
           "## Public rtb_2026_industry, cells that differ", ""]
    if pub_lines:
        md += ["| Industry | Column | Ours | Public | Abs diff |", "|---|---|---:|---:|---:|"] + pub_lines
    else:
        md.append("None.")
    md += ["", "## Public data_sec_agg (feeds shortrunseries C7:C13 and Q7:Q13)", "",
           "| Year | n ours | n public | worth ours | worth public | |",
           "|---|---:|---:|---:|---:|---|"] + dsa_lines

    out = sql_out_dir() / "check_rtb_ca.md"
    out.write_text("\n".join(md) + "\n", encoding="utf-8")

    print("Query 1 check")
    for c in checks:
        print(f"  {'PASS' if c['ok'] else 'FAIL'}  {c['name']:<50} n={c['n']:<5} maxdiff={_fmt(c['maxdiff'])}")
    print(f"{failures} failing check(s). Report: {out}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
