"""Helpers shared by the step-1 checkers (py/check_*.py).

Each checker compares the tables a query exports (data-raw/sql-out/py/) with
the authors' own outputs in the confidential bundle, and the R exports with
the Python ones. A check is a dict with keys name, ok, n (cells or rows
compared), maxdiff and note. Reports print counts and maximum differences
only, never rows of the private files.
"""
import csv
import math
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from bundle_paths import sql_out_dir  # noqa: E402

csv.field_size_limit(10_000_000)

TOL_PARITY = 1e-9


def missing(v):
    return (v is None or v == "" or v == "NA" or v == "#N/A"
            or (isinstance(v, float) and math.isnan(v)))


def num(v):
    if missing(v):
        return None
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, datetime):
        raise ValueError("date in a numeric column")
    return float(v)


def text(v):
    if missing(v):
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def read_csv_dicts(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def read_export(table, lang="py"):
    return read_csv_dicts(sql_out_dir(lang) / f"{table}.csv")


def read_xlsx_dicts(path, sheet=None, header_row=1):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(min_row=header_row, values_only=True)
    header = list(next(rows))
    keep = [i for i, h in enumerate(header) if h is not None]
    out = []
    for r in rows:
        if all(v is None for v in r):
            continue
        out.append({header[i]: (r[i] if i < len(r) else None) for i in keep})
    wb.close()
    return out


def num_diff(a, b, rel=False):
    """|a - b| (or relative to max(1, |b|)); inf when one side is missing."""
    a, b = num(a), num(b)
    if a is None and b is None:
        return 0.0
    if (a is None) != (b is None):
        return math.inf
    d = abs(a - b)
    return d / max(1.0, abs(b)) if rel else d


def compare_rows(name, ours, key, num_cols, text_cols, tol, rename=None, rel=False):
    """Compare two equally ordered lists of dicts cell by cell.

    rename maps our column names to the key's. Returns a check dict plus the
    per-column count of cells off (for the report).
    """
    rename = rename or {}
    off = {c: 0 for c in num_cols + text_cols}
    maxd = 0.0
    for o, k in zip(ours, key):
        for c in num_cols:
            d = num_diff(o[c], k[rename.get(c, c)], rel)
            if d > tol:
                off[c] += 1
            if d != math.inf:
                maxd = max(maxd, d)
        for c in text_cols:
            if text(o[c]) != text(k[rename.get(c, c)]):
                off[c] += 1
    bad = {c: v for c, v in off.items() if v}
    ok = len(ours) == len(key) and not bad
    note = (f"rows {len(ours)} vs {len(key)}; cells off: "
            + (", ".join(f"{c} {v}" for c, v in bad.items()) or "none"))
    return {"name": name, "ok": ok, "n": min(len(ours), len(key)), "maxdiff": maxd,
            "note": note, "off": off}


def compare_keyed(name, ours, key, key_cols, num_cols, text_cols, tol, rename=None, rel=False):
    """Align two lists of dicts on key_cols (same names on both sides after rename)."""
    rename = rename or {}

    def k_of(r, side):
        return tuple(text(r[c if side == "o" else rename.get(c, c)]) for c in key_cols)

    o = {k_of(r, "o"): r for r in ours}
    k = {k_of(r, "k"): r for r in key}
    common = [x for x in o if x in k]
    chk = compare_rows(name, [o[x] for x in common], [k[x] for x in common],
                       num_cols, text_cols, tol, rename, rel)
    only_o, only_k = len(set(o) - set(k)), len(set(k) - set(o))
    dup = (len(ours) - len(o)) + (len(key) - len(k))
    chk["ok"] = chk["ok"] and only_o == 0 and only_k == 0 and dup == 0
    chk["note"] = (f"keys {len(o)} vs {len(k)}; only in ours {only_o}, only in key {only_k}"
                   + (f"; duplicate keys {dup}" if dup else "") + "; " + chk["note"])
    chk["only_ours"] = sorted(set(o) - set(k))
    chk["only_key"] = sorted(set(k) - set(o))
    return chk


def parity(tables):
    """R export vs Python export, every cell, numbers within TOL_PARITY."""
    out = []
    for t in tables:
        p, r = sql_out_dir("py") / f"{t}.csv", sql_out_dir("r") / f"{t}.csv"
        if not r.exists():
            out.append({"name": f"R vs Python: {t}", "ok": False, "n": 0, "maxdiff": math.nan,
                        "note": "R export missing; run Rscript R/run_sql.R first"})
            continue
        a, b = read_csv_dicts(p), read_csv_dicts(r)
        same_cols = (list(a[0]) == list(b[0])) if a and b else len(a) == len(b)
        maxd, text_bad = 0.0, 0
        for x, y in zip(a, b):
            for c in x:
                u, v = x[c], y.get(c)
                try:
                    d = abs(float(u) - float(v))
                    maxd = max(maxd, d / max(1.0, abs(float(v))))
                except (TypeError, ValueError):
                    text_bad += u != v
        ok = same_cols and len(a) == len(b) and maxd <= TOL_PARITY and text_bad == 0
        out.append({"name": f"R vs Python: {t}", "ok": ok, "n": len(a), "maxdiff": maxd,
                    "note": f"rows {len(a)} vs {len(b)}; text cells off {text_bad}; "
                            "max diff relative to max(1, |r|)"})
    return out


def fmt(x):
    if x == math.inf:
        return "inf"
    if isinstance(x, float) and math.isnan(x):
        return "n/a"
    return f"{x:.3g}"


def write_report(title, checks, report_name, extra_md=None, tol_line=""):
    """Write data-raw/sql-out/<report_name>.md and print a summary; return the exit code."""
    failures = sum(not c["ok"] for c in checks)
    md = [f"# {title}", "",
          f"Generated {datetime.now():%Y-%m-%d %H:%M}. Counts, totals and maximum differences",
          "only; no row of the private files is printed.", ""]
    if tol_line:
        md += [tol_line, ""]
    md += ["## Summary", "", "| Check | Compared | Max diff | Status | Detail |",
           "|---|---:|---:|---|---|"]
    for c in checks:
        md.append(f"| {c['name']} | {c['n']} | {fmt(c['maxdiff'])} | "
                  f"{'PASS' if c['ok'] else 'FAIL'} | {c['note']} |")
    md += [""] + (extra_md or [])
    out = sql_out_dir() / f"{report_name}.md"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(md) + "\n", encoding="utf-8")
    print(title)
    for c in checks:
        print(f"  {'PASS' if c['ok'] else 'FAIL'}  {c['name']:<58} n={c['n']:<7} "
              f"maxdiff={fmt(c['maxdiff'])}")
        if not c["ok"]:
            print(f"        {c['note']}")
    print(f"{failures} failing check(s). Report: {out}")
    return 1 if failures else 0
